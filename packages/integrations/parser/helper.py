"""Private stdin-only parser. No document paths, URLs, logging or network I/O."""
import hashlib
import io
import json
import resource
import sys
import zipfile

MAX_INPUT = 5 * 1024 * 1024
MAX_OUTPUT = 64000


def extract(data, kind, ocr=False, tessdata="/models", with_coverage=False):
    complete = True
    source = kind
    if ocr:
        for name, expected in [("rus", "e16e5e036cce1d9ec2b00063cf8b54472625b9e14d893a169e2b0dedeb4df225"), ("eng", "7d4322bd2a7749724879683fc3912cb542f19906c83bcc1a52132556427170b2")]:
            with open(tessdata+"/"+name+".traineddata", "rb") as model:
                if hashlib.sha256(model.read(5000000)).hexdigest() != expected:
                    return None
    if kind in ("jpeg", "png"):
        import pymupdf
        with pymupdf.open(stream=data, filetype=kind) as image:
            if image[0].rect.width * image[0].rect.height > 4000000:
                return None
            data = image.convert_to_pdf()
        kind = "pdf"
    if not data or len(data) > MAX_INPUT:
        return None
    if kind == "pdf":
        import pymupdf
        with pymupdf.open(stream=data, filetype="pdf") as document:
            if document.needs_pass or len(document) > 40:
                return None
            parts = []
            total = 0
            ocr_pages = 0
            pixels = 0
            for page in document:
                text = page.get_text()
                if page.get_images() and (not ocr or len(text.strip()) >= 50):
                    complete = False
                if len(text.strip()) < 50 and not ocr:
                    complete = False
                if len(text.strip()) < 50 and ocr:
                    ocr_pages += 1
                    page_pixels = int(page.rect.width * 150/72) * int(page.rect.height * 150/72)
                    pixels += page_pixels
                    if ocr_pages > 5 or page_pixels > 4000000 or pixels > 16000000:
                        return None
                    text = page.get_text(textpage=page.get_textpage_ocr(language="rus+eng", dpi=150, full=True, tessdata=tessdata))
                    if len(text.strip()) < 50:
                        return None
                total += len(text.encode("utf-8"))
                if total > MAX_OUTPUT:
                    return None
                parts.append(text)
            text = "\n".join(parts)
    elif kind == "docx":
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            items = archive.infolist()
            if any(i.filename.startswith("word/media/") for i in items):
                complete = False
            if len(items) > 256 or sum(i.file_size for i in items) > 20 * 1024 * 1024:
                return None
            if any(i.flag_bits & 1 or i.file_size > 10 * 1024 * 1024 or
                   i.file_size > max(1, i.compress_size) * 200 for i in items):
                return None
            # Block external entities/relationships before the XML library sees them.
            for item in items:
                if item.filename.endswith((".xml", ".rels")):
                    # Parse bounded XML with explicit non-resolving defaults. Attribute
                    # syntax/escaping/encoding cannot disguise an External relationship.
                    from lxml import etree
                    parser = etree.XMLParser(resolve_entities=False, no_network=True,
                                             load_dtd=False, huge_tree=False)
                    try:
                        tree = etree.fromstring(archive.read(item), parser)
                    except etree.XMLSyntaxError:
                        return None
                    if tree.getroottree().docinfo.doctype:
                        return None
                    if item.filename.endswith(".rels") and any(
                        node.get("TargetMode", "").strip().casefold() == "external"
                        for node in tree.iter()
                    ):
                        return None
        from docx import Document
        document = Document(io.BytesIO(data))
        text = "\n".join([p.text for p in document.paragraphs] +
                         [cell.text for table in document.tables for row in table.rows for cell in row.cells])
    else:
        return None
    text = text.strip()
    if len(text) < 50 or len(text.encode("utf-8")) > MAX_OUTPUT:
        return None
    result = {"source": source, "content": text}
    if with_coverage:
        result["coverage"] = "complete" if complete else "partial"
    return result


def main():
    resource.setrlimit(resource.RLIMIT_CORE, (0, 0))
    resource.setrlimit(resource.RLIMIT_CPU, (3, 3))
    resource.setrlimit(resource.RLIMIT_AS, (384 * 1024 * 1024, 384 * 1024 * 1024))
    resource.setrlimit(resource.RLIMIT_FSIZE, (0, 0))
    resource.setrlimit(resource.RLIMIT_NOFILE, (32, 32))
    try:
        data = sys.stdin.buffer.read(MAX_INPUT + 1)
        args = sys.argv[2:]
        kind = sys.argv[1] if 2 <= len(sys.argv) <= 4 and set(args) <= {"ocr", "coverage"} else ""
        result = extract(data, kind) if not args else extract(data, kind, ocr="ocr" in args, with_coverage="coverage" in args)
        sys.stdout.write(json.dumps(result, ensure_ascii=True))
    except Exception:
        sys.stdout.write("null")


if __name__ == "__main__":
    main()
