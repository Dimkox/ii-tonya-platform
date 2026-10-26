import io
import zipfile
import pymupdf
from docx import Document
from helper import extract

TEXT = "A private synthetic professional biography used for deterministic parser acceptance. " * 3

def test_real_text_pdf():
    document = pymupdf.open()
    page = document.new_page()
    page.insert_text((50, 50), TEXT)
    assert len(extract(document.tobytes(), "pdf")["content"]) >= 50


def test_real_docx():
    document = Document()
    document.add_paragraph(TEXT)
    data = io.BytesIO()
    document.save(data)
    assert extract(data.getvalue(), "docx")["content"] == TEXT.strip()


def test_empty_scan_encrypted_and_page_limit():
    document = pymupdf.open()
    document.new_page()
    assert extract(document.tobytes(), "pdf") is None
    assert extract(document.tobytes(encryption=pymupdf.PDF_ENCRYPT_AES_256, owner_pw="secret", user_pw="secret"), "pdf") is None
    for _ in range(40):
        document.new_page()
    assert extract(document.tobytes(), "pdf") is None


def test_expansion_and_xml_entities_rejected():
    data = io.BytesIO()
    with zipfile.ZipFile(data, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("word/document.xml", "x" * 100000)
    assert extract(data.getvalue(), "docx") is None
    data = io.BytesIO()
    with zipfile.ZipFile(data, "w") as archive:
        archive.writestr("word/document.xml", '<!DOCTYPE a [<!ENTITY x SYSTEM "file:///private">]>')
    assert extract(data.getvalue(), "docx") is None


def test_short_text():
    document = Document()
    document.add_paragraph("short")
    data = io.BytesIO()
    document.save(data)
    assert extract(data.getvalue(), "docx") is None

# Real subprocess evidence calls main(), so its limits affect only this child.
import json
import subprocess
import sys
from pathlib import Path
import pytest


def main_probe(body, cwd):
    source = f'''import importlib.util,sys,json,resource,signal,os
spec=importlib.util.spec_from_file_location("private_helper",{str(Path(__file__).with_name("helper.py"))!r})
helper=importlib.util.module_from_spec(spec);spec.loader.exec_module(helper)
def probe(data,kind):
{body}
helper.extract=probe
sys.argv=["helper","pdf"]
helper.main()
'''
    return subprocess.run([sys.executable, "-I", "-c", source], input=b"synthetic", capture_output=True, timeout=8, cwd=cwd)


def test_main_resource_limits_and_empty_file_boundary(tmp_path):
    result=main_probe('''    limits={name:list(resource.getrlimit(getattr(resource,name))) for name in ["RLIMIT_CPU","RLIMIT_AS","RLIMIT_NOFILE","RLIMIT_CORE","RLIMIT_FSIZE"]}
    open("empty-proof","xb").close()
    signal.signal(signal.SIGXFSZ,signal.SIG_IGN)
    try:
        with open("empty-proof","wb") as target: target.write(b"x")
        limits["growth_prevented"]=False
    except OSError: limits["growth_prevented"]=True
    limits["empty_file_exists"]=os.path.exists("empty-proof")
    handles=[]
    try:
        for _ in range(40): handles.append(open("/dev/null","rb"))
    except OSError: limits["fd_bound_enforced"]=True
    finally:
        for handle in handles: handle.close()
    return limits''',tmp_path)
    assert result.returncode==0 and result.stderr==b""
    actual=json.loads(result.stdout)
    assert actual["RLIMIT_CPU"]==[3,3]
    assert actual["RLIMIT_AS"]==[384*1024*1024]*2
    assert actual["RLIMIT_NOFILE"]==[32,32]
    assert actual["RLIMIT_CORE"]==[0,0]
    assert actual["RLIMIT_FSIZE"]==[0,0]
    assert actual["empty_file_exists"] and actual["growth_prevented"] and actual["fd_bound_enforced"]
    assert (tmp_path/"empty-proof").stat().st_size==0


def test_main_memory_failure_has_fixed_output(tmp_path):
    result=main_probe('    return bytearray(400*1024*1024)',tmp_path)
    assert result.returncode==0 and result.stdout==b"null" and result.stderr==b""


def test_main_cpu_bound_terminates_only_child(tmp_path):
    result=main_probe('    while True: pass',tmp_path)
    assert result.returncode in (-9,-24) and result.stdout==b"" and result.stderr==b""


@pytest.mark.parametrize("attribute,encoding", [
    ("TargetMode='External'","utf-8"),
    ('TargetMode = "External"',"utf-8"),
    ('TargetMode="Ext&#101;rnal"',"utf-8"),
    ('TargetMode="External"',"utf-16"),
])
def test_equivalent_external_relationship_attributes(attribute,encoding):
    document=Document();document.add_paragraph(TEXT);source=io.BytesIO();document.save(source)
    target=io.BytesIO()
    with zipfile.ZipFile(source) as original,zipfile.ZipFile(target,"w") as updated:
        for item in original.infolist():
            data=original.read(item)
            if item.filename=="word/_rels/document.xml.rels":
                xml=data.decode("utf-8").replace("</Relationships>",f'<Relationship Id="syntheticExternal" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink" Target="https://example.invalid/private" {attribute}/></Relationships>')
                if encoding=="utf-16":xml=xml.replace("UTF-8","UTF-16")
                data=xml.encode(encoding)
            updated.writestr(item,data)
    assert extract(target.getvalue(),"docx") is None


@pytest.mark.skipif(not __import__("os").environ.get("II_TONYA_OCR_TEST_MODELS"), reason="optional pinned OCR model fixture absent")
def test_real_cyrillic_scan_ocr():
    document=pymupdf.open();page=document.new_page(width=600,height=400)
    font=pymupdf.Font(fontfile='/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf')
    writer=pymupdf.TextWriter(page.rect)
    writer.append((30,70),'Специальность: Психолог',font=font,fontsize=24)
    writer.append((30,120),'Учебное заведение: Университет',font=font,fontsize=24)
    writer.append((30,170),'Год: 2020',font=font,fontsize=24)
    writer.write_text(page)
    pix=page.get_pixmap(dpi=150)
    scanned=pymupdf.open();scan=scanned.new_page(width=600,height=400)
    scan.insert_image(scan.rect,stream=pix.tobytes('png'))
    result=extract(scanned.tobytes(),'pdf',ocr=True,tessdata=__import__('os').environ['II_TONYA_OCR_TEST_MODELS'])
    assert result and 'Психолог' in result['content'] and '2020' in result['content']
